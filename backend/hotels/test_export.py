from io import BytesIO
from unittest.mock import patch

from django.test import SimpleTestCase
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from hotels.export_api import MAX_EXPORT_ROWS
from hotels.services.export.excel import build_hotel_workbook, export_columns


CONTEXT = {
    'location': 'Delhi Airport',
    'latitude': 28.5562,
    'longitude': 77.1,
    'radius': 5000,
    'provider': 'OpenStreetMap',
}

HOTEL = {
    'name': 'Roseate House',
    'distance_km': 2.1,
    'address': 'Aerocity, New Delhi',
    'phone': '+911171558800',
    'email': 'info@roseatehotels.com',
    'website': 'https://www.roseatehotels.com/',
    'brand': 'Roseate Hotels',
    'stars': '5',
    'latitude': 28.55,
    'longitude': 77.12,
    'source': 'OpenStreetMap',
    'enrichment_status': 'FOUND',
    'enrichment_sources': {
        'address': 'OpenStreetMap',
        'phone': 'https://www.roseatehotels.com/contact',
        'email': 'https://www.roseatehotels.com/contact',
        'website': 'OpenStreetMap',
        'brand': 'OpenStreetMap',
    },
    'manager_contacts': [{
        'name': 'Manager Name',
        'title': 'General Manager',
        'department': 'General Management',
        'email': 'manager@roseatehotels.com',
        'phone': '+911112345678',
        'linkedin_url': 'https://linkedin.com/in/manager',
    }],
}


class ExcelWorkbookTests(SimpleTestCase):
    def workbook(self, hotels=None):
        data, _ = build_hotel_workbook(hotels or [HOTEL], CONTEXT)
        return load_workbook(BytesIO(data))

    def test_workbook_contains_search_summary(self):
        self.assertIn('Search Summary', self.workbook().sheetnames)

    def test_workbook_contains_hotels(self):
        self.assertIn('Businesses', self.workbook().sheetnames)

    def test_workbook_contains_structured_enrichment_sheets(self):
        hotel = {
            **HOTEL,
            'business_emails': [{'email': 'sales@hotel.example', 'type': 'sales',
                                 'source_url': 'https://hotel.example/contact'}],
            'business_phones': [{'phone': '+1 212 555 0100', 'type': 'general',
                                 'source_url': 'https://hotel.example/contact'}],
            'decision_makers': [{'name': 'Jane Doe', 'title': 'IT Director',
                                 'department': 'Information Technology',
                                 'role_group': 'it_leadership',
                                 'source_url': 'https://hotel.example/team'}],
            'social_profiles': {'linkedin': 'https://linkedin.com/company/hotel'},
            'social_profile_sources': {'linkedin': 'https://hotel.example/'},
        }
        workbook = self.workbook([hotel])
        self.assertIn('Business Contacts', workbook.sheetnames)
        self.assertIn('Decision Makers', workbook.sheetnames)
        self.assertIn('Social Profiles', workbook.sheetnames)
        self.assertEqual(workbook['Business Contacts'].max_row, 3)
        self.assertEqual(workbook['Decision Makers'].cell(2, 2).value, 'Jane Doe')
        self.assertEqual(
            workbook['Social Profiles'].cell(2, 2).value,
            'https://linkedin.com/company/hotel',
        )

    def test_headers_are_correct(self):
        sheet = self.workbook()['Businesses']
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            [label for label, _ in export_columns()],
        )

    def test_hotel_row_values(self):
        row = self.workbook()['Businesses'][2]
        self.assertEqual(row[1].value, 'Roseate House')
        self.assertEqual(row[2].value, 2.1)

    def test_phone_plus_prefix_is_preserved_safely(self):
        phone = self.workbook()['Businesses'].cell(2, 5)
        self.assertEqual(phone.value, "'+911171558800")
        self.assertEqual(phone.number_format, '@')

    def test_unicode_hotel_name_is_preserved(self):
        hotel = {**HOTEL, 'name': 'होटल स्वागत'}
        self.assertEqual(self.workbook([hotel])['Businesses'].cell(2, 2).value, 'होटल स्वागत')

    def test_website_hyperlink(self):
        cell = self.workbook()['Businesses'].cell(2, 7)
        self.assertEqual(cell.hyperlink.target, HOTEL['website'])

    def test_manager_columns(self):
        sheet = self.workbook()['Businesses']
        self.assertEqual(sheet.cell(2, 14).value, 'Manager Name')
        self.assertEqual(sheet.cell(2, 18).hyperlink.target, 'https://linkedin.com/in/manager')

    def test_field_source_columns(self):
        sheet = self.workbook()['Businesses']
        headers = [cell.value for cell in sheet[1]]
        phone_source_column = headers.index('Phone Source') + 1
        self.assertEqual(
            sheet.cell(2, phone_source_column).value,
            'https://www.roseatehotels.com/contact',
        )

    def test_missing_values_are_blank(self):
        hotel = {'name': 'Minimal Hotel'}
        sheet = self.workbook([hotel])['Businesses']
        self.assertIsNone(sheet.cell(2, 4).value)
        self.assertIsNone(sheet.cell(2, 14).value)

    def test_formula_injection_is_escaped(self):
        hotel = {**HOTEL, 'name': '=HYPERLINK("https://bad.example")'}
        value = self.workbook([hotel])['Businesses'].cell(2, 2).value
        self.assertTrue(value.startswith("'="))


class ExcelExportAPITests(APITestCase):
    url = '/api/hotels/export/excel/'

    def payload(self, hotels=None):
        return {'hotels': [HOTEL] if hotels is None else hotels, 'context': CONTEXT}

    def test_excel_export_success(self):
        response = self.client.post(self.url, self.payload(), format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('business-leads-delhi-airport-', response['Content-Disposition'])

    def test_invalid_hotels_payload(self):
        response = self.client.post(
            self.url, {'hotels': 'not-a-list', 'context': CONTEXT}, format='json'
        )
        self.assertEqual(response.status_code, 400)

    def test_empty_list_handling(self):
        response = self.client.post(self.url, self.payload([]), format='json')
        self.assertEqual(response.status_code, 400)

    def test_row_limit_validation(self):
        hotels = [{'name': f'Hotel {index}'} for index in range(MAX_EXPORT_ROWS + 1)]
        response = self.client.post(self.url, self.payload(hotels), format='json')
        self.assertEqual(response.status_code, 400)

    @patch('requests.get')
    @patch('requests.post')
    def test_export_does_not_call_external_apis(self, mock_post, mock_get):
        response = self.client.post(self.url, self.payload(), format='json')
        self.assertEqual(response.status_code, 200)
        mock_get.assert_not_called()
        mock_post.assert_not_called()


class ExportPreparationAPITests(APITestCase):
    url = '/api/hotels/export/prepare/'

    @patch('hotels.services.export.enrichment.enrich_hotels_bulk')
    def test_website_business_is_automatically_enriched_with_official_data(self, enrich):
        enrich.return_value = {'results': [{
            'hotel_name': 'Hotel', 'status': 'FOUND',
            'hotel': {'website': 'https://hotel.example', 'email': 'info@hotel.example'},
            'sources': {'email': 'https://hotel.example/contact'},
            'business_emails': [{'email': 'info@hotel.example', 'type': 'general',
                                 'source_url': 'https://hotel.example/contact'}],
            'business_phones': [{'phone': '+1 212 555 0100', 'type': 'general'}],
            'social_profiles': {'linkedin': 'https://linkedin.com/company/hotel'},
            'social_profile_sources': {'linkedin': 'https://hotel.example'},
            'decision_makers': [
                {'name': 'Jane Doe', 'title': 'IT Director',
                 'role_group': 'it_leadership', 'source': 'Official Website'},
            ],
        }]}
        response = self.client.post(self.url, {'hotels': [
            {'name': 'Hotel', 'website': 'https://hotel.example'},
        ]}, format='json')
        self.assertEqual(response.status_code, 200)
        hotel = response.json()['hotels'][0]
        self.assertEqual(hotel['business_emails'][0]['email'], 'info@hotel.example')
        self.assertEqual(hotel['decision_makers'][0]['title'], 'IT Director')
        enrich.assert_called_once()

    @patch('hotels.services.export.enrichment.enrich_hotels_bulk')
    def test_no_website_and_existing_enrichment_are_not_fetched(self, enrich):
        ready = {
            'name': 'Ready', 'website': 'https://ready.example',
            'enrichment_status': 'NOT_FOUND', 'business_emails': [],
            'business_phones': [], 'social_profiles': {}, 'decision_makers': [],
        }
        response = self.client.post(self.url, {'hotels': [
            {'name': 'No Website', 'address': 'Delhi'}, ready,
        ]}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()['hotels']), 2)
        enrich.assert_not_called()

    @patch('hotels.services.export.enrichment.enrich_hotels_bulk')
    def test_more_than_ten_are_processed_in_bounded_batches(self, enrich):
        def response_for(batch):
            return {'results': [{
                'hotel_name': hotel['name'], 'status': 'NOT_FOUND',
                'hotel': hotel, 'sources': {},
            } for hotel in batch]}
        enrich.side_effect = response_for
        hotels = [
            {'name': f'Hotel {index}', 'website': f'https://hotel-{index}.example'}
            for index in range(23)
        ]
        response = self.client.post(self.url, {'hotels': hotels}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual([len(call.args[0]) for call in enrich.call_args_list], [10, 10, 3])

    @patch('hotels.services.export.enrichment.enrich_hotels_bulk')
    def test_individual_failure_is_preserved_and_does_not_fail_export(self, enrich):
        enrich.return_value = {'results': [
            {'hotel_name': 'Bad', 'status': 'ERROR', 'hotel': {'website': 'https://bad.example'}},
            {'hotel_name': 'Good', 'status': 'FOUND', 'hotel': {
                'website': 'https://good.example', 'email': 'info@good.example'},
             'business_emails': [{'email': 'info@good.example'}]},
        ]}
        response = self.client.post(self.url, {'hotels': [
            {'name': 'Bad', 'website': 'https://bad.example', 'address': 'Original'},
            {'name': 'Good', 'website': 'https://good.example'},
        ]}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['hotels'][0]['address'], 'Original')
        self.assertEqual(response.json()['hotels'][1]['email'], 'info@good.example')

    @patch('hotels.services.people_enrichment.apollo.search_decision_makers')
    @patch('hotels.services.people_enrichment.zoominfo.ZoomInfoPeopleEnrichmentProvider.search_decision_makers')
    @patch('hotels.services.export.enrichment.enrich_hotels_bulk')
    def test_paid_people_providers_are_never_called(self, enrich, zoominfo, apollo):
        enrich.return_value = {'results': [{
            'hotel_name': 'Hotel', 'status': 'NOT_FOUND',
            'hotel': {'website': 'https://hotel.example'},
        }]}
        response = self.client.post(self.url, {'hotels': [
            {'name': 'Hotel', 'website': 'https://hotel.example'},
        ]}, format='json')
        self.assertEqual(response.status_code, 200)
        apollo.assert_not_called()
        zoominfo.assert_not_called()
