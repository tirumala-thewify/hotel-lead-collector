from unittest.mock import patch

from django.test import SimpleTestCase
from openpyxl import load_workbook

from hotels.services.enrichment.hotel_website import (
    ContactPageParser, _extract_whatsapp_contacts, _whatsapp_url_contact,
    enrich_hotel_from_website,
)
from hotels.services.export.excel import build_hotel_workbook


PUBLIC_DNS = [(2, 1, 6, '', ('93.184.216.34', 443))]


def response(html):
    class FakeResponse:
        status_code = 200
        is_redirect = False
        is_permanent_redirect = False
        encoding = 'utf-8'
        headers = {'Content-Type': 'text/html'}

        def iter_content(self, chunk_size):
            yield html.encode()
    return FakeResponse()


class WhatsAppExtractionTests(SimpleTestCase):
    def parse(self, html):
        parser = ContactPageParser()
        parser.feed(html)
        return _extract_whatsapp_contacts(parser)

    def test_supported_urls_and_query_strings(self):
        for url, kind in (
            ('https://wa.me/919876543210', 'wa.me'),
            ('https://wa.me/919876543210?text=Hello', 'wa.me'),
            ('https://api.whatsapp.com/send?phone=919876543210', 'api.whatsapp.com'),
            ('https://api.whatsapp.com/send/?phone=919876543210&text=Hi', 'api.whatsapp.com'),
        ):
            with self.subTest(url=url):
                item = _whatsapp_url_contact(url)
                self.assertEqual(item['normalized'], '+919876543210')
                self.assertEqual(item['evidence_type'], kind)

    def test_visible_label_and_labelled_tel_anchor(self):
        contacts = self.parse(
            '<p>WhatsApp: +91 98765 43210</p>'
            '<a href="tel:+44 20 1234 5678">WhatsApp us</a>'
        )
        self.assertEqual(
            {item['evidence_type'] for item in contacts},
            {'whatsapp_labelled_text', 'whatsapp_labelled_link'},
        )

    def test_ordinary_phone_and_mobile_are_not_evidence(self):
        self.assertEqual(self.parse('<p>Phone: +91 98765 43210 Mobile: +91 99999 99999</p>'), [])

    def test_fake_hosts_and_unsafe_urls_are_rejected(self):
        for url in (
            'https://wa.me.example.com/919876543210',
            'https://whatsapp.com.fake-domain.com/send?phone=919876543210',
            'https://api.whatsapp.com.example.org/send?phone=919876543210',
            'javascript:alert(1)', 'data:text/plain,919876543210',
            'https://user:pass@wa.me/919876543210',
        ):
            with self.subTest(url=url):
                self.assertIsNone(_whatsapp_url_contact(url))

    def test_no_country_code_is_invented_and_suspicious_value_is_preserved(self):
        item = self.parse('<p>WhatsApp Number: 987 654 321</p>')[0]
        self.assertEqual(item['number'], '987 654 321')
        self.assertEqual(item['normalized'], '987654321')
        self.assertIsNone(_whatsapp_url_contact('https://wa.me/12-34'))


class WhatsAppCrawlTests(SimpleTestCase):
    @patch('hotels.services.enrichment.hotel_website.socket.getaddrinfo', return_value=PUBLIC_DNS)
    @patch('hotels.services.enrichment.hotel_website.requests.get')
    def test_uses_fetched_pages_deduplicates_and_preserves_strongest_source(self, get, _dns):
        home = ('<a href="/contact">Contact</a><p>Phone: +91 11111 11111</p>'
                '<p>WhatsApp: +91 98765 43210</p>')
        contact = ('<a href="https://wa.me/919876543210?text=Hi">Chat</a>'
                   '<a href="https://wa.me/919999999999">Chat</a>')
        get.side_effect = [response(home), response(contact)]
        result = enrich_hotel_from_website({'name': 'Hotel', 'website': 'https://hotel.example/'})
        self.assertEqual(len(result['whatsapp_contacts']), 2)
        first = result['whatsapp_contacts'][0]
        self.assertEqual(first['evidence_type'], 'wa.me')
        self.assertEqual(first['source_url'], 'https://hotel.example/contact')
        self.assertEqual(result['business_phones'][0]['normalized'], '+911111111111')
        self.assertEqual(get.call_count, 2)  # official pages only; wa.me was never fetched

    def test_workbook_only_contains_confirmed_contacts(self):
        hotels = [{
            'name': 'Hotel', 'website': 'https://hotel.example', 'phone': '+1 111 111 1111',
            'whatsapp_contacts': [{
                'number': '+1 222 222 2222', 'normalized': '+12222222222',
                'status': 'CONFIRMED_PUBLIC', 'evidence_type': 'wa.me',
                'source_url': 'https://hotel.example/contact',
            }],
        }, {'name': 'Phone Only', 'phone': '+1 333 333 3333'}]
        content, _ = build_hotel_workbook(hotels, {
            'location': 'Test', 'latitude': 1, 'longitude': 2,
            'radius': 3, 'provider': 'Browser Search',
        })
        from io import BytesIO
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(
            workbook.sheetnames,
            ['Search Summary', 'Businesses', 'Business Contacts', 'Decision Makers',
             'Social Profiles', 'WhatsApp Contacts'],
        )
        sheet = workbook['WhatsApp Contacts']
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet.cell(2, 3).value, "'+12222222222")
