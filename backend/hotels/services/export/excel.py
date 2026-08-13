from io import BytesIO
from urllib.parse import urlparse

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MANAGER_DEPARTMENTS = {
    'General Management': 'General Manager',
    'Marketing': 'Marketing Manager',
    'Information Technology': 'IT Manager',
}


def export_columns():
    columns = [
        ('S.No', 'serial'),
        ('Hotel Name', 'name'),
        ('Distance from Search Location (km)', 'distance_km'),
        ('Address', 'address'),
        ('Phone', 'phone'),
        ('Email', 'email'),
        ('Website', 'website'),
        ('Brand', 'brand'),
        ('Stars', 'stars'),
        ('Latitude', 'latitude'),
        ('Longitude', 'longitude'),
        ('Source', 'source'),
        ('Enrichment Status', 'enrichment_status'),
    ]
    for label in ('General Manager', 'Marketing Manager', 'IT Manager'):
        for field in ('Name', 'Title', 'Email', 'Phone', 'LinkedIn'):
            columns.append((f'{label} {field}', f'{label}:{field.lower()}'))
    for field in ('Address', 'Phone', 'Email', 'Website', 'Brand'):
        columns.append((f'{field} Source', f'source:{field.lower()}'))
    columns.append(('Rating', 'rating'))
    return columns


def sanitize_spreadsheet_value(value):
    if value is None:
        return ''
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return f"'{value}"
    return value


def _manager_values(hotel):
    values = {}
    contacts = hotel.get('manager_contacts') or []
    for contact in contacts:
        if not isinstance(contact, dict):
            continue
        label = MANAGER_DEPARTMENTS.get(contact.get('department'))
        if not label or f'{label}:name' in values:
            continue
        values.update({
            f'{label}:name': contact.get('name'),
            f'{label}:title': contact.get('title'),
            f'{label}:email': contact.get('email'),
            f'{label}:phone': contact.get('phone'),
            f'{label}:linkedin': contact.get('linkedin_url'),
        })
    return values


def hotel_export_row(hotel, index):
    source_values = {
        f'source:{field}': (hotel.get('enrichment_sources') or {}).get(field)
        for field in ('address', 'phone', 'email', 'website', 'brand')
    }
    values = {**hotel, **_manager_values(hotel), **source_values, 'serial': index}
    return [sanitize_spreadsheet_value(values.get(key)) for _, key in export_columns()]


def _safe_hyperlink(value):
    if not isinstance(value, str):
        return None
    try:
        parsed = urlparse(value)
    except ValueError:
        return None
    return value if parsed.scheme in {'http', 'https'} and parsed.netloc else None


def _add_structured_sheet(workbook, title, headers, rows, link_columns=()):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='172033')
    for row in rows:
        sheet.append([sanitize_spreadsheet_value(value) for value in row])
        for column in link_columns:
            cell = sheet.cell(sheet.max_row, column)
            link = _safe_hyperlink(cell.value)
            if link:
                cell.hyperlink = link
                cell.style = 'Hyperlink'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{sheet.max_row}'
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(18, min(50, len(header) + 4))
    return sheet


def build_hotel_workbook(hotels, context):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Search Summary'
    exported_at = timezone.now()
    summary_rows = [
        ('Location', context.get('location')),
        ('Latitude', context.get('latitude')),
        ('Longitude', context.get('longitude')),
        ('Radius', context.get('radius')),
        ('Hotel Provider', context.get('provider')),
        ('Total Hotels', len(hotels)),
        ('Exported At', exported_at.isoformat()),
    ]
    for label, value in summary_rows:
        summary.append([label, sanitize_spreadsheet_value(value)])
    for cell in summary['A']:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='246BFD')
    summary.column_dimensions['A'].width = 22
    summary.column_dimensions['B'].width = 60

    sheet = workbook.create_sheet('Hotels')
    headers = [label for label, _ in export_columns()]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='172033')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for index, hotel in enumerate(hotels, start=1):
        sheet.append(hotel_export_row(hotel, index))
        row_number = sheet.max_row
        if index % 2 == 0:
            for cell in sheet[row_number]:
                cell.fill = PatternFill('solid', fgColor='F4F7FB')
        for cell in sheet[row_number]:
            cell.alignment = Alignment(vertical='top', wrap_text=True)

        website_cell = sheet.cell(row_number, 7)
        website_link = _safe_hyperlink(hotel.get('website'))
        if website_link:
            website_cell.hyperlink = website_link
            website_cell.style = 'Hyperlink'
        email_cell = sheet.cell(row_number, 6)
        if hotel.get('email'):
            email_cell.hyperlink = f"mailto:{hotel['email']}"
            email_cell.style = 'Hyperlink'
        for column in (5, 17, 22, 27):
            sheet.cell(row_number, column).number_format = '@'
        for column in (18, 23, 28):
            link = _safe_hyperlink(sheet.cell(row_number, column).value)
            if link:
                sheet.cell(row_number, column).hyperlink = link
                sheet.cell(row_number, column).style = 'Hyperlink'
        distance = sheet.cell(row_number, 3)
        if isinstance(distance.value, (int, float)):
            distance.number_format = '0.00'

    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{sheet.max_row}'
    widths = [8, 28, 18, 42, 18, 28, 30, 20, 10, 14, 14, 18, 18] + [24] * 20
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 34

    business_rows = []
    people_rows = []
    social_rows = []
    whatsapp_rows = []
    for hotel in hotels:
        for item in hotel.get('business_emails') or []:
            business_rows.append((hotel.get('name'), item.get('type'),
                                  item.get('email'), None, item.get('source_url')))
        for item in hotel.get('business_phones') or []:
            business_rows.append((hotel.get('name'), item.get('type'), None,
                                  item.get('phone'), item.get('source_url')))
        for contact in hotel.get('decision_makers') or []:
            people_rows.append((
                hotel.get('name'), contact.get('name'), contact.get('title'),
                contact.get('role_group'),
                contact.get('business_email') or contact.get('email'),
                contact.get('phone'), contact.get('source'), contact.get('source_url'),
                contact.get('confidence'),
            ))
        profiles = hotel.get('social_profiles') or {}
        for contact in hotel.get('whatsapp_contacts') or []:
            if contact.get('status') != 'CONFIRMED_PUBLIC':
                continue
            whatsapp_rows.append((
                hotel.get('name'), contact.get('number'), contact.get('normalized'),
                contact.get('status'), contact.get('evidence_type'),
                contact.get('source_url'), hotel.get('website'), hotel.get('phone'),
            ))
        profile_sources = hotel.get('social_profile_sources') or {}
        source_urls = list(dict.fromkeys(
            profile_sources.get(platform) for platform in ('linkedin', 'facebook', 'instagram')
            if profile_sources.get(platform)
        ))
        social_rows.append((
            hotel.get('name'), profiles.get('linkedin'), profiles.get('facebook'),
            profiles.get('instagram'), '; '.join(source_urls),
        ))
    _add_structured_sheet(
        workbook, 'Business Contacts',
        ['Business Name', 'Contact Type', 'Email', 'Phone', 'Source URL'],
        business_rows, link_columns=(5,),
    )
    _add_structured_sheet(
        workbook, 'Decision Makers',
        ['Business Name', 'Person Name', 'Title', 'Role Group', 'Email',
         'Phone', 'Source', 'Source URL', 'Confidence'],
        people_rows, link_columns=(8,),
    )
    _add_structured_sheet(
        workbook, 'Social Profiles',
        ['Business Name', 'LinkedIn', 'Facebook', 'Instagram', 'Source URL'],
        social_rows, link_columns=(2, 3, 4, 5),
    )
    _add_structured_sheet(
        workbook, 'WhatsApp Contacts',
        ['Hotel Name', 'WhatsApp Number', 'Normalized WhatsApp Number',
         'WhatsApp Status', 'Evidence Type', 'Source URL', 'Website', 'Business Phone'],
        whatsapp_rows, link_columns=(6, 7),
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue(), exported_at
